import sys, json, re, subprocess, warnings, os
warnings.filterwarnings("ignore"); sys.path.insert(0, "src")
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import parse_xml
from docx.text.paragraph import Paragraph
from rb_helpers import *
from parser import parse_pcap

TEMPLATE = "/mnt/user-data/uploads/CNEH_CBP_FORMAT.docx"
TITLE = "NETWORK TRAFFIC ANALYSIS & PROTOCOL-LEVEL ANOMALY DETECTION SYSTEM"
SUBJECT = "COMPUTER NETWORKS AND ETHICAL HACKING (22PC1CY201)"
R = json.load(open("results/results.json")); BASE = json.load(open("results/baseline.json"))
W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

# ---- facts measured from the held-out normal capture (for Table 2.6 / 4.1) ----
dn = parse_pcap("pcaps/normal_test.pcap")
ICMP_MAX = int(dn[(dn.proto == "ICMP") & (dn.icmp_type == 8)].groupby("src").dst.nunique().max())
DHCP_SERVERS = int(dn[(dn.proto == "DHCP") & (dn.dhcp == 2)].smac.nunique())
PROTO = dn.proto.value_counts()
NSRC = int(dn.src.nunique()); SYNPS = float((dn["flags"] == "S").sum() / BASE["dur"])
console = open("results/console_combined.txt").read().splitlines()
ALERTS = []
for ln in console:
    m = re.match(r"\[(\w+)\s*\] t=\s*([\d.]+)s (.+?)\s+src=(.+?)\s+\| (.*)", ln)
    if m: ALERTS.append(m.groups())
assert len(ALERTS) == 8, ALERTS
ml = R["ml"]; sens = R["sensitivity"]
LAT = {k: R[k]["latency"] for k in ["portscan", "pingsweep", "arpspoof", "dnsenum", "ftpbrute", "synflood", "roguedhcp"]}
assert all(R[k]["detected"] for k in LAT) and R["normal"]["alerts"] == 0 and R["normal_test"]["alerts"] == 0

def build(pm, out):
    doc = Document(TEMPLATE); ctx = Ctx(doc); body = doc.element.body; ch = list(body.iterchildren())
    # bullet numbering (black bullets)
    numbering = doc.part.numbering_part.element
    absn = parse_xml(f'<w:abstractNum {W_NS} w:abstractNumId="90"><w:multiLevelType w:val="hybridMultilevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="•"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman" w:hint="default"/><w:color w:val="000000"/></w:rPr></w:lvl>'
        '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="–"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="1080" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman" w:hint="default"/><w:color w:val="000000"/></w:rPr></w:lvl></w:abstractNum>')
    fn = numbering.find(qn("w:num")); fn.addprevious(absn) if fn is not None else numbering.append(absn)
    numbering.append(parse_xml(f'<w:num {W_NS} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))
    P_ = lambda *a, **k: P(ctx, *a, **k); B_ = lambda *a, **k: bullet(ctx, *a, **k)
    pg = lambda key: str(pm.get(key, "–"))

    # ============ FRONT MATTER: fill placeholders ============
    replace_in_para(ch[0], "TITLE OF THE PROJECT", TITLE)
    replace_in_para(ch[3], "NAME OF LABORATORY (SUBJECT CODE).", SUBJECT)
    replace_in_para(ch[29], "TITLE OF THE COURSE-BASED PROJECT", TITLE)
    replace_in_para(ch[45], "COURSE BASED PROJECT TITLE", TITLE)

    # Acknowledgement
    ack = [
     "We express our sincere gratitude to everyone who supported us during the course-based project “Network Traffic Analysis & Protocol-Level Anomaly Detection System”.",
     "We are deeply thankful to our project supervisor, [[NAME OF THE SUPERVISOR]], [[Designation of the Supervisor]], for the guidance, technical support and continuous monitoring that kept this work on track from the first lab design to the final report.",
     "We thank **Dr. P V Siva Kumar**, Associate Professor & I/C Head of the Department of CSE-(CyS, DS) and AI&DS, for providing the facilities, encouragement and administrative support needed to carry out the project.",
     "We are grateful to the faculty members and lab staff of the Computer Networks and Ethical Hacking course, whose lectures and laboratory sessions on network protocols, footprinting, scanning and enumeration formed the technical foundation of this work.",
     "We thank **Vallurupalli Nageswara Rao Vignana Jyothi Institute of Engineering & Technology** for the infrastructure, computing resources and learning environment, and the open-source communities behind Scapy, Wireshark, Nmap, scikit-learn and Streamlit, whose tools made this project possible.",
     "Finally, we thank our families and classmates for their support and feedback.",
    ]
    for t in ack:
        p = P_(t, "j", 12, after=8); ch[71].addprevious(p._p)
    for i in range(64, 71): body.remove(ch[i])
    Paragraph(ch[62], doc._body).paragraph_format.page_break_before = True
    for i in list(range(55, 62)) + list(range(76, 85)): body.remove(ch[i])
    Paragraph(ch[85], doc._body).paragraph_format.page_break_before = True
    for i in range(86, 126): body.remove(ch[i])
    anchor = ch[126]
    anchor_pPr = anchor.find(qn("w:pPr")); sp = parse_xml(f'<w:spacing {W_NS} w:before="0" w:after="0" w:line="20" w:lineRule="exact"/>'); anchor_pPr.insert(0, sp)
    for i in range(127, 336): body.remove(ch[i])

    # ============ BODY ============
    # ---------- Chapter 1 ----------
    chapter(ctx, 1, "INTRODUCTION")
    heading(ctx, "1.1 Background")
    P_("Every packet that crosses a network follows rules written into protocol specifications. A TCP connection opens with a three-way handshake {r:rfc793}; an ARP request is answered by the owner of the requested address {r:rfc826}; a DHCP client receives an offer from the server that manages the address pool {r:rfc2131}; an FTP server answers a failed login with reply code 530 {r:rfc959}. Normal applications obey these exchanges, whereas most reconnaissance and attack tools break them in characteristic ways: a port scanner sends many SYN segments and never completes a handshake, an ARP spoofer announces an address that belongs to someone else, and a brute-forcing tool produces long runs of failed logins.")
    P_("The Computer Networks and Ethical Hacking course follows the same path an attacker takes. Units I–III describe the link, network, transport and application layers and the protocols that run on them {r:forouzan}{r:kurose}; Units IV and V cover footprinting, scanning with Nmap, enumeration and vulnerability analysis {r:engebretson}{r:lyon}. This project looks at the same material from the defender’s side: if we know how a protocol behaves when it is used correctly, we can recognise the moment it is abused.")
    heading(ctx, "1.2 Need and Relevance of the Work")
    P_("Local area networks are built on protocols that were designed for cooperation rather than distrust. ARP and DHCP have no built-in authentication, DNS servers are sometimes configured to hand out whole zones on request, and FTP sends credentials in clear text. Reconnaissance such as port scanning and ping sweeps usually precedes a real attack, so noticing it early gives administrators time to react. Production intrusion-detection systems such as Snort {r:roesch} and Bro/Zeek {r:paxson} solve this problem at scale, but their rule languages and internal architecture are large, and for a student they behave as black boxes.")
    P_("A compact, readable detector has educational and practical value. It shows exactly which header field or reply code identifies each attack, it makes thresholds explicit so they can be tuned and criticised, and it can be run on any packet capture without special infrastructure. Building and evaluating such a detector also exercises the whole syllabus: lab set-up (Weeks 4–7), scanning and enumeration (Weeks 8–11) and protocol analysis (Units I–III).")
    heading(ctx, "1.3 Problem Statement")
    P_("To design and implement a software system that analyses captured network traffic at the protocol level, characterises normal behaviour, and automatically detects and reports reconnaissance and attack activity — port scanning, ping sweeps, ARP spoofing, DNS enumeration and zone transfer, FTP brute force, SYN flooding and rogue DHCP servers — with low detection delay and few false alarms.")
    heading(ctx, "1.4 Objectives of the Project")
    for t in [
        "**O1.** Design a reproducible lab topology and a capture workflow for normal and attack traffic.",
        "**O2.** Characterise normal traffic (protocol mix, packet rate, connection behaviour) to obtain a baseline for threshold selection.",
        "**O3.** Implement protocol-level detectors for at least five attack classes covered in the syllabus.",
        "**O4.** Measure detection rate and detection latency for every attack class.",
        "**O5.** Keep false alarms on normal traffic as low as possible and justify every threshold.",
        "**O6.** Present results through a command-line tool and an interactive dashboard.",
        "**O7.** Explore an unsupervised machine-learning detector (Isolation Forest) and compare it with the rule-based approach."]:
        B_(t)
    heading(ctx, "1.5 Scope of the Work")
    P_("The work covers offline analysis of IPv4 Ethernet traffic stored in .pcap files; the same files can be produced by tcpdump or Wireshark {r:wireshark} on a live network. Seven attack classes are implemented as rule-based detectors, and one unsupervised model is evaluated alongside them. Results are reported through a command-line tool and a Streamlit dashboard {r:streamlit}.")
    P_("The following are outside the scope: inspection of encrypted payloads, IPv6 and wireless protocols, inline blocking (intrusion prevention) and line-rate performance on production links. The experimental datasets were generated with Scapy {r:scapy} to reproduce the on-the-wire signature of each attack tool (Section 2.4), so the reported figures describe the behaviour of the detectors on controlled, fully labelled data and should not be read as a statement about live production networks.")
    heading(ctx, "1.6 Relation to the Course Syllabus")
    table(ctx, ["Syllabus part", "Course outcome", "Used in this project"], [
        ["Unit I – Data link layer, addressing, switching", "CO-1", "Ethernet frame parsing; source-MAC evidence for ARP spoofing and rogue DHCP; host-only switch in the lab"],
        ["Unit II – Network layer: IPv4, ARP, DHCP, ICMP", "CO-2", "Ping-sweep, ARP-spoofing and rogue-DHCP detectors; IP header analysis"],
        ["Unit III – Transport (TCP/UDP) and application layer (DNS, FTP, HTTP)", "CO-3", "TCP flag analysis for port-scan and SYN-flood detection; DNS and FTP protocol detectors"],
        ["Unit IV – Footprinting and scanning (Nmap)", "CO-4", "Host-discovery and port-scan attack scenarios; scan detectors"],
        ["Unit V – Enumeration, vulnerability scanning", "CO-5", "DNS enumeration and zone-transfer detection"],
        ["Lab Weeks 4–7 – networking commands, FTP server, network set-up, VirtualBox and Kali", "Lab CO-2, CO-3", "Lab topology, FTP service, attacker and victim machines"],
        ["Lab Weeks 8–11 – information gathering, DNS/mail enumeration, Nmap", "Lab CO-4, CO-5", "Attack scenarios whose traffic the system must recognise"]],
        [2900, 1100, 4300], 10, ["l", "c", "l"], caption=("T", "1.1", "Mapping of syllabus topics to project components"))
    heading(ctx, "1.7 Organisation of the Report")
    P_("Chapter 2 explains the approach: architecture, lab, datasets, parser, detection algorithms and the machine-learning layer. Chapter 3 shows test cases and the actual output of the tool. Chapter 4 reports detection results, false-alarm analysis, threshold sensitivity, the comparison with machine learning and validation against the objectives. Chapter 5 concludes with limitations and future scope. References and appendices (team contribution, source code and real-lab capture commands) follow.")

    # ---------- Chapter 2 ----------
    chapter(ctx, 2, "APPROACH / METHODOLOGY")
    heading(ctx, "2.1 Overall Approach")
    P_("The project follows a five-stage pipeline: **capture → parse → detect → alert → evaluate**. Traffic is stored as .pcap files so that every experiment is repeatable; the parser turns each packet into one row of a table; detectors are functions that scan this table for protocol behaviour that should never, or only rarely, occur in normal traffic; alerts carry the evidence that triggered them; and an evaluation script measures detection, latency and false alarms against ground truth.")
    P_("A **rule-based** design was chosen as the core because protocol exchanges are deterministic: a SYN that is never followed by an ACK, an IP address that suddenly maps to a second MAC address, or a 530 reply repeated five times are unambiguous facts, and an alert can state exactly why it fired. Threshold values are taken from measurements of the normal baseline rather than guessed. An **Isolation Forest** {r:liu} is added as an exploratory second opinion, to test whether a model trained only on normal traffic can find the same attacks without protocol knowledge.")
    heading(ctx, "2.2 System Architecture")
    figure(ctx, "figs/arch.png", 5.7, "2.1", "System architecture of the anomaly detection pipeline")
    table(ctx, ["Module", "Source file", "Responsibility"], [
        ["Dataset generator", "src/gen_pcaps.py", "Builds normal and attack .pcap files with Scapy (fixed random seed)"],
        ["Parser", "src/parser.py", "Reads a .pcap and returns a pandas table, one row per packet"],
        ["Detectors", "src/detectors.py", "Seven rule-based detectors and run_all() driver"],
        ["ML layer", "src/ml_iforest.py", "Window features and Isolation Forest training / prediction"],
        ["CLI", "src/run_detector.py", "Prints alerts with severity, time, source and evidence; optional CSV"],
        ["Evaluation", "src/evaluate.py", "Computes detection, latency, false alerts, sensitivity, ML comparison"],
        ["Dashboard", "dashboard/app.py", "Streamlit page: metrics, protocol chart, timeline, alert table"]],
        [1700, 2000, 4600], 10, caption=("T", "2.1", "System modules and source files"))
    heading(ctx, "2.3 Lab Environment")
    P_("The reference lab is a VirtualBox host-only network, 192.168.56.0/24, shown in Fig. 2.2. A Kali Linux machine is the attacker (192.168.56.66), an Ubuntu server is the victim (192.168.56.10) running FTP, HTTP and DNS services, ten client addresses (192.168.56.20–29) generate everyday traffic, and a monitor machine captures frames in promiscuous mode. A host-only network keeps the attack traffic away from any real network. This is the same environment built in Lab Weeks 5–7 of the course.")
    figure(ctx, "figs/topology.png", 5.6, "2.2", "Lab network topology (192.168.56.0/24)")
    table(ctx, ["Component", "Tool / version", "Role"], [
        ["Hypervisor and lab OS", "VirtualBox; Kali Linux; Ubuntu Server", "Reference lab for live capture (Appendix C)"],
        ["Traffic capture / inspection", "tcpdump, Wireshark {r:wireshark}", "Create and inspect .pcap files"],
        ["Attack tools being modelled", "Nmap {r:lyon}, Hydra, hping3, arpspoof, dig, dnsmasq", "Source of the attack signatures that are detected"],
        ["Language and runtime", "Python 3.12.3", "All project code"],
        ["Packet parsing and generation", "Scapy 2.8.0 {r:scapy}", "Read .pcap files; synthesise datasets"],
        ["Data handling", "pandas 3.0.2 {r:pandas}, NumPy 2.4.4", "Packet table and window statistics"],
        ["Machine learning", "scikit-learn 1.8.0 {r:sklearn}", "Isolation Forest"],
        ["Visualisation", "matplotlib 3.10.8, Plotly, Streamlit 1.66.0 {r:streamlit}", "Figures and dashboard"]],
        [2300, 3300, 2700], 10, caption=("T", "2.2", "Software and tools"))
    heading(ctx, "2.4 Dataset Generation")
    P_("Two kinds of traffic are needed: normal traffic to learn the baseline and to measure false alarms, and attack traffic with known ground truth. Both were produced by the script gen_pcaps.py using Scapy, with a fixed random seed so that every run produces identical files. This approach makes the experiments exactly repeatable and every packet fully labelled; captures taken from the VMs of the reference lab can be analysed by the same tools without change (Appendix C).")
    P_("**Normal traffic model.** Ten clients generate events as a Poisson process of about six events per second. Each event is one of: a DNS query and answer (35 %), an HTTP session (30 %), an HTTPS-like TLS session (13 %), a successful FTP login (10 %), an ICMP echo and reply (5 %), an ARP request and reply (4 %) or a DHCP discover and offer (3 %). TCP sessions include the complete SYN, SYN-ACK, ACK handshake, a data exchange and a FIN exchange.")
    P_("**Attack models.** Each attack reproduces the packet-level behaviour of the real tool (Table 2.3). Attack captures consist of 120 seconds of normal background traffic into which the attack is injected at t = 40 s, so detectors must find the attack among legitimate traffic. A combined capture contains all seven attacks one after another at 40-second intervals. Table 2.4 lists the resulting files.")
    table(ctx, ["Attack", "Real tool and typical command", "Signature reproduced in the capture"], [
        ["Port scan", "nmap -sS -p 1-10000 192.168.56.10", "1,000 SYNs to different ports from one source at about 250 per second; RST (closed) or SYN-ACK (open) replies"],
        ["Ping sweep", "nmap -sn 192.168.56.0/24", "ICMP echo requests to 254 different hosts, one every 20 ms"],
        ["ARP spoofing", "arpspoof -t 192.168.56.20 192.168.56.1", "Unsolicited ARP replies claiming the gateway IP with the attacker’s MAC, repeated every second"],
        ["DNS enumeration / zone transfer", "dnsrecon -t brt; dig axfr lab.local", "60 queries for different host names at 20 per second, followed by one AXFR request over TCP/53"],
        ["FTP brute force", "hydra -l admin -P words.txt ftp://192.168.56.10", "60 FTP sessions from one source, each ending with “530 Login incorrect”"],
        ["SYN flood", "hping3 -S --flood --rand-source -p 80", "3,000 SYNs in 3 s to one port from random spoofed sources; no handshake completes"],
        ["Rogue DHCP", "dnsmasq --dhcp-range=…", "DHCPOFFER messages from a second server MAC"]],
        [1500, 2900, 3900], 9.5, caption=("T", "2.3", "Attack scenarios, reference tools and the signatures reproduced"))
    ds = [["normal.pcap", R["normal"]["packets"], f"{R['normal']['duration']:.0f}", "Baseline normal traffic (training of the ML model)"],
          ["normal_test.pcap", R["normal_test"]["packets"], f"{R['normal_test']['duration']:.0f}", "Held-out normal traffic (different seed) for false-alarm tests"]]
    for k, lab in [("portscan", "Port scan"), ("pingsweep", "Ping sweep"), ("arpspoof", "ARP spoofing"), ("dnsenum", "DNS enumeration + AXFR"), ("ftpbrute", "FTP brute force"), ("synflood", "SYN flood"), ("roguedhcp", "Rogue DHCP")]:
        ds.append([f"{k}.pcap", R[k]["packets"], "120", f"{lab} injected at t = 40 s into normal traffic"])
    ds.append(["slowscan.pcap", R["slowscan_rule"]["packets"], "150", "Slow port scan: 80 ports, one SYN every 1.5 s (evasion test)"])
    ds.append(["combined.pcap", R["combined"]["packets"], "300", "All seven attacks in sequence (t = 20, 60, …, 260 s)"])
    table(ctx, ["File", "Packets", "Duration (s)", "Content"], [[a, f"{b:,}", c, d] for a, b, c, d in ds], [1900, 1000, 1100, 4300], 9.5, ["l", "r", "r", "l"], caption=("T", "2.4", "Datasets used in the experiments (pcaps/ folder)"))
    heading(ctx, "2.5 Packet Parser")
    P_("The parser (parser.py) reads a capture packet by packet with Scapy’s PcapReader {r:scapy} and produces a pandas DataFrame {r:pandas} with one row per packet. Keeping the whole capture in one table allows every detector to be written as a few lines of grouping and window logic. Application-layer evidence that detectors need (FTP reply code, DNS query name and type, DHCP message type) is extracted once during parsing.")
    table(ctx, ["Group", "Fields extracted"], [
        ["Frame", "timestamp, frame length, source MAC"],
        ["Network", "source IP, destination IP; for ARP the sender IP/MAC and opcode"],
        ["Transport", "protocol label, source and destination port, TCP flag string (S, SA, A, PA, FA, RA)"],
        ["ICMP", "message type (8 = echo request, 0 = echo reply)"],
        ["DNS", "query name, query type, AXFR flag (TCP/53 request with QTYPE 252 {r:rfc5936})"],
        ["FTP", "flag for server reply code 530 (login failed) {r:rfc959}"],
        ["DHCP", "message type (discover, offer, …) and the MAC of the server that sent an offer {r:rfc2131}"]],
        [1500, 6800], 10, caption=("T", "2.5", "Fields extracted by the parser"))
    heading(ctx, "2.6 Detection Algorithms")
    P_("Most detectors share one primitive: a **sliding window that counts distinct items**. Events are processed in time order; items older than the window length W are discarded; an alert is raised the first time the number of distinct items inside the window reaches the threshold K. The pseudo-code below is the core of the port-scan, ping-sweep and DNS-enumeration detectors.")
    code(ctx, ["function SLIDING_DISTINCT(events, W, K):",
               "    window <- empty queue ; count <- empty map",
               "    for (t, item) in events sorted by time:",
               "        push (t, item) to window ; count[item] += 1",
               "        while t - window.front.time > W:",
               "            (t0, old) <- pop window ; count[old] -= 1",
               "            if count[old] = 0: delete count[old]",
               "        if size(count) >= K: return (t, size(count))   # alert",
               "    return none"])
    dets = [
      ("2.6.1 Port Scan (Unit IV, Nmap)", "A TCP port scan sends SYN segments to many different ports and abandons the connection when a port answers {r:lyon}{r:rfc793}. Normal clients talk to a few well-known ports. **Rule:** for each (source, destination) pair, count distinct destination ports of packets with only the SYN flag set; alert when K ≥ 20 within W = 5 s. Sequential hypothesis testing {r:jung} is a more adaptive alternative, discussed under future scope."),
      ("2.6.2 Ping Sweep", "Host discovery sends ICMP echo requests {r:rfc792} to every address of a subnet. **Rule:** per source, count distinct destination addresses of echo requests (type 8); alert when K ≥ 10 within W = 10 s."),
      ("2.6.3 ARP Spoofing (Unit II)", "ARP has no authentication {r:rfc826}; an attacker poisons caches by sending replies that bind a victim’s IP address to the attacker’s MAC address. **Rule:** maintain a table IP → MAC learned from ARP replies; alert when a reply binds an already-known IP to a different MAC. The alert names the new (offending) MAC address."),
      ("2.6.4 DNS Enumeration and Zone Transfer (Unit V)", "Enumeration tools guess host names in bulk, and a misconfigured server may answer an AXFR request with its entire zone {r:rfc1035}{r:rfc5936}. **Rule 1:** per source, count distinct query names; alert at K ≥ 30 within W = 10 s. **Rule 2:** alert on any TCP/53 request with QTYPE 252 (AXFR)."),
      ("2.6.5 FTP Brute Force (Week 5 service)", "FTP signals a failed login with reply code 530 {r:rfc959}. **Rule:** group 530 replies by the client that receives them; alert when five or more occur within 60 s."),
      ("2.6.6 SYN Flood", "A SYN flood exhausts the server’s half-open connection table by sending SYNs that are never completed {r:rfc4987}. Unlike a port scan, all SYNs target the same port. **Rule:** for each (destination, port) pair, alert when 100 or more SYN segments arrive within 1 s; the alert also reports how many distinct sources were involved, which exposes spoofed addresses."),
      ("2.6.7 Rogue DHCP Server", "A client accepts the first offer it receives, so a second DHCP server can redirect traffic through the attacker {r:rfc2131}. **Rule:** collect the MAC addresses that send DHCPOFFER messages; if more than one exists, every MAC other than the most frequent (the legitimate server) is reported."),
    ]
    for h, t in dets:
        heading(ctx, h, 3); P_(t)
    P_("Table 2.6 lists the thresholds together with the largest value observed in the held-out normal capture. The ratio between the two is the safety margin against false alarms.")
    table(ctx, ["Detector", "Parameter (rule)", "Threshold", "Normal maximum"], [
        ["Port scan", "distinct ports, one source → one host, 5 s", "20", "3"],
        ["Ping sweep", "distinct hosts pinged by one source, 10 s", "10", f"{ICMP_MAX} (whole capture)"],
        ["ARP spoofing", "changes of an IP → MAC binding", "1", "0"],
        ["DNS enumeration", "distinct query names, one source, 10 s", "30", "5"],
        ["DNS zone transfer", "AXFR requests", "1", "0"],
        ["FTP brute force", "530 replies to one client, 60 s", "5", "0"],
        ["SYN flood", "SYNs per second to one host:port", "100", "8"],
        ["Rogue DHCP", "distinct DHCP-offer server MACs", "2", f"{DHCP_SERVERS}"]],
        [1900, 3400, 1200, 1800], 10, ["l", "l", "c", "c"], caption=("T", "2.6", "Detector thresholds and observed normal behaviour (normal_test.pcap)"))
    heading(ctx, "2.7 Isolation Forest (Exploratory ML Layer)")
    P_("Isolation Forest {r:liu} isolates observations with random splits; points that are isolated after few splits are anomalies. The model is trained only on windows from normal.pcap (200 trees, contamination 0.01, fixed seed) using the scikit-learn implementation {r:sklearn}. For every source address and every 5-second window, six features are computed (Table 2.7). No attack data is used for training, so the model has to recognise attacks purely as deviations from normal behaviour.")
    table(ctx, ["Feature", "Meaning"], [
        ["pkts", "packets sent by the source in the window"],
        ["bytes_avg", "mean packet length"],
        ["uniq_dports", "distinct destination ports contacted"],
        ["uniq_dsts", "distinct destination addresses contacted"],
        ["syn_ratio", "fraction of packets that are bare SYN segments"],
        ["icmp_ratio", "fraction of packets that are ICMP"]],
        [2000, 6300], 10, caption=("T", "2.7", "Window features for the Isolation Forest"))
    heading(ctx, "2.8 Alerting, Command-Line Tool and Dashboard")
    P_("Every detector returns alert records containing the **type, time, source, evidence text and severity** (MEDIUM, HIGH or CRITICAL). The command-line tool run_detector.py parses a capture, runs all detectors and prints one line per alert; with --csv the alerts are also saved. The Streamlit dashboard (dashboard/app.py) accepts an uploaded .pcap or a sample file and shows four summary metrics, a protocol-distribution chart, a packets-per-second timeline with a dashed marker for each alert, and a sortable alert table.")

    # ---------- Chapter 3 ----------
    chapter(ctx, 3, "TEST CASES / OUTPUT")
    heading(ctx, "3.1 Test Procedure")
    P_("Each test case runs the detectors on one capture and compares the alerts with the expected result. A test passes when the expected alert type is raised against the attacker and no alert is raised against any legitimate host. The tests are executed with the commands below; evaluate.py runs all of them and stores the results in results/results.json.")
    code(ctx, ["python src/gen_pcaps.py                       # build all datasets in pcaps/",
               "python src/run_detector.py pcaps/portscan.pcap   # run one test case",
               "python src/evaluate.py                         # run every test case, write results/",
               "streamlit run dashboard/app.py                 # open the dashboard"], 8.5)
    heading(ctx, "3.2 Test Cases and Outcomes")
    tc = [["TC-01", "Normal baseline", "normal.pcap", "No alert", "0 alerts", "Pass"],
          ["TC-02", "Held-out normal, 10 min", "normal_test.pcap", "No alert", "0 alerts", "Pass"]]
    for i, (k, lab, exp) in enumerate([("portscan", "Port scan", "Port Scan"), ("pingsweep", "Ping sweep", "Ping Sweep"), ("arpspoof", "ARP spoofing", "ARP Spoofing"),
            ("dnsenum", "DNS enumeration + AXFR", "DNS Enumeration, DNS Zone Transfer"), ("ftpbrute", "FTP brute force", "FTP Brute Force"), ("synflood", "SYN flood", "SYN Flood"), ("roguedhcp", "Rogue DHCP server", "Rogue DHCP")]):
        tc.append([f"TC-{i+3:02d}", lab, f"{k}.pcap", exp + " alert(s) on attacker", f"{R[k]['alerts']} alert(s), first at +{LAT[k]:.2f} s, {R[k]['false_alerts']} false", "Pass"])
    tc.append(["TC-10", "All attacks in one capture", "combined.pcap", "8 alerts, all on attacker", f"{R['combined']['alerts']} alerts, {R['combined']['attacker_alerts']} on attacker", "Pass"])
    tc.append(["TC-11", "Slow port scan (evasion)", "slowscan.pcap", "Port Scan alert", "No alert", "Fail (limitation)"])
    table(ctx, ["ID", "Scenario", "Input", "Expected", "Observed", "Result"], tc, [800, 1550, 1450, 1700, 1950, 850], 9, ["c", "l", "l", "l", "l", "c"], caption=("T", "3.1", "Test cases and outcomes"))
    heading(ctx, "3.3 Command-Line Output")
    P_("Fig. 3.1 shows the output of the tool on the combined capture, and Fig. 3.2 on 10 minutes of normal traffic. In the first, each of the eight alerts names the attacker (192.168.56.66, its MAC address 08:00:27:ee:00:66, or “100 sources” for the spoofed SYN flood) and the evidence; in the second, no alert is raised. Times are measured from the first packet of the capture.")
    figure(ctx, "figs/console_combined.png", 5.8, "3.1", "Console output of run_detector.py on combined.pcap")
    figure(ctx, "figs/console_normal_test.png", 5.8, "3.2", "Console output of run_detector.py on normal_test.pcap")
    table(ctx, ["Time (s)", "Alert type", "Severity", "Source", "Evidence"],
          [[a[1], a[2], a[0], a[3], a[4]] for a in ALERTS], [750, 1500, 950, 1550, 3550], 8.5, ["r", "l", "l", "l", "l"], caption=("T", "3.2", "Alert records produced for the combined capture"))
    heading(ctx, "3.4 Dashboard Output")
    P_("The dashboard presents the same alerts graphically. Its timeline panel (Fig. 3.3) plots packets per second on a logarithmic axis and draws a dashed line at each alert. The bursts at the port scan, ping sweep, FTP brute force and SYN flood are visible in the traffic itself, while the ARP spoofing, DNS and rogue-DHCP alerts correspond to very small traffic changes that a volume-based monitor would miss. The protocol chart for normal traffic is shown in Fig. 4.1.")
    figure(ctx, "figs/timeline.png", 5.7, "3.3", "Traffic timeline of combined.pcap with alert markers (dashboard view)")

    # ---------- Chapter 4 ----------
    chapter(ctx, 4, "RESULTS, ANALYSIS AND VALIDATION")
    heading(ctx, "4.1 Evaluation Method")
    P_("An attack is counted as **detected** when every expected alert type is raised and the alert names the attacker. **Detection latency** is the time between the first attack packet (t = 40 s in the single-attack captures) and the first alert. A **false alert** is any alert that does not name the attacker. False alerts are counted on the baseline capture, on the 10-minute held-out normal capture, and on the legitimate background traffic inside every attack capture. For the Isolation Forest, a 5-second source window is flagged when the model marks it anomalous, and an attack counts as found when at least one of the attacker’s windows is flagged.")
    heading(ctx, "4.2 Baseline Traffic Characteristics")
    P_(f"The held-out normal capture contains {BASE['n']:,} packets over {BASE['dur']:.0f} seconds, an average of {BASE['pps']:.1f} packets per second with a mean packet length of {BASE['avg']:.0f} bytes, from {NSRC} distinct source addresses. Bare SYN segments occur at only {SYNPS:.1f} per second. Fig. 4.1 and Table 4.1 show the protocol mix; HTTP, TLS, FTP and DNS together account for about 95 % of packets, whereas ARP, ICMP and DHCP are rare.")
    figure(ctx, "figs/proto_pie.png", 5.0, "4.1", "Protocol distribution of the baseline traffic (normal_test.pcap)")
    table(ctx, ["Protocol", "Packets", "Share (%)"], [[k, f"{v:,}", f"{v/PROTO.sum()*100:.1f}"] for k, v in PROTO.items()] + [["Total", f"{PROTO.sum():,}", "100.0"]],
          [3000, 2500, 2500], 10, ["l", "r", "r"], caption=("T", "4.1", "Baseline traffic characteristics: packets per protocol"))
    P_("The baseline matters because it makes the thresholds defensible. For example, a normal host contacted at most three distinct ports within 5 seconds, so a threshold of 20 gives a margin of almost seven times, and the highest SYN rate towards a single port was 8 per second against a threshold of 100.")
    heading(ctx, "4.3 Detection Results")
    rows = []
    for k, lab, exp in [("portscan", "Port scan", "Port Scan"), ("pingsweep", "Ping sweep", "Ping Sweep"), ("arpspoof", "ARP spoofing", "ARP Spoofing"), ("dnsenum", "DNS enumeration + AXFR", "DNS Enumeration, DNS Zone Transfer"),
                        ("ftpbrute", "FTP brute force", "FTP Brute Force"), ("synflood", "SYN flood", "SYN Flood"), ("roguedhcp", "Rogue DHCP", "Rogue DHCP")]:
        rows.append([lab, f"{R[k]['packets']:,}", exp, "Yes", f"{LAT[k]:.2f}", R[k]["alerts"], R[k]["false_alerts"]])
    table(ctx, ["Scenario", "Packets", "Expected alert(s)", "Detected", "Latency (s)", "Alerts", "False alerts"], rows, [1600, 900, 2250, 900, 900, 850, 900], 9.5, ["l", "r", "l", "c", "r", "c", "c"],
          caption=("T", "4.2", "Detection results for the seven attack scenarios"))
    P_("All seven attack scenarios, and all eight alert types, were detected: a detection rate of 7 of 7 with no false alerts in any attack capture. Fig. 4.2 shows the detection latency. It is explained by arithmetic: a detector alerts as soon as its threshold is reached, so the delay is roughly the threshold divided by the attacker’s rate. The port scan needs 20 ports at one probe per 4 ms (0.08 s), the ping sweep needs 10 hosts at 20 ms (0.18 s), the SYN flood needs 100 SYNs at 1 ms (0.10 s), and DNS enumeration needs 30 names at 50 ms (1.45 s). ARP spoofing and rogue DHCP are decided from a single packet and are therefore reported immediately.")
    figure(ctx, "figs/latency.png", 5.0, "4.2", "Detection latency per attack scenario")
    heading(ctx, "4.4 False-Alarm Analysis")
    P_(f"The rule-based detectors raised **no alert** on normal.pcap ({R['normal']['packets']:,} packets) or on the 10-minute normal_test.pcap ({R['normal_test']['packets']:,} packets), and no alert in the 120 seconds of background traffic contained in each attack capture. Comparing Table 2.6 shows why: normal behaviour stays far below every threshold (3 against 20 ports, 5 against 30 DNS names, 8 against 100 SYNs per second), and the remaining detectors look for events that never occur in normal traffic (a second MAC for one IP, a 530 reply, an AXFR request, a second DHCP server). The Isolation Forest, in contrast, flagged {ml['normal_test_flagged']} of {ml['normal_test_windows']:,} normal windows ({ml['normal_test_flagged']/ml['normal_test_windows']*100:.1f} %).")
    heading(ctx, "4.5 Threshold Sensitivity")
    P_("To check how much the port-scan result depends on the chosen threshold, the detector was run with thresholds from 2 to 300 distinct ports (Table 4.3, Fig. 4.3) on the port-scan capture, the normal capture and the slow-scan capture.")
    table(ctx, ["Threshold (ports / 5 s)", "Port scan detected", "False alerts (normal, 10 min)", "Slow scan detected"],
          [[s["threshold"], "Yes" if s["scan_detected"] else "No", s["fp_on_normal"], "Yes" if s["slow_detected"] else "No"] for s in sens], [2100, 1900, 2400, 1900], 10, ["c"] * 4,
          caption=("T", "4.3", "Port-scan threshold sensitivity"))
    figure(ctx, "figs/sensitivity.png", 5.2, "4.3", "Effect of the port-scan threshold on false alerts and slow-scan detection")
    P_("The fast scan is found at every threshold, so the result is robust. Thresholds of 2 and 3 ports cause 10 false alerts because ordinary clients occasionally open two or three services within 5 seconds, while any value from 5 to 300 gives none. The slow scan, which probes one port every 1.5 s (at most four ports in any 5-second window), is found only by the thresholds that also produce false alarms. A fixed-window rule therefore cannot separate a slow scan from normal behaviour; this trade-off is the main limitation of the approach (Section 4.7).")
    heading(ctx, "4.6 Rule-Based Detection versus Isolation Forest")
    mlrows = []
    lab = {"portscan": "Port scan", "pingsweep": "Ping sweep", "arpspoof": "ARP spoofing", "dnsenum": "DNS enumeration + AXFR", "ftpbrute": "FTP brute force", "synflood": "SYN flood", "roguedhcp": "Rogue DHCP"}
    for k in lab:
        m = ml[k]; mlrows.append([lab[k], m["att_windows"], m["att_flagged"], "Found" if m["att_flagged"] else "Missed", "Found"])
    mlrows.append(["Slow port scan", ml["slowscan"]["att_windows"], ml["slowscan"]["att_flagged"], "Missed", "Missed"])
    table(ctx, ["Scenario", "Attacker windows", "Flagged by model", "Isolation Forest", "Rule-based"], mlrows, [2400, 1500, 1500, 1500, 1400], 10, ["l", "r", "r", "c", "c"],
          caption=("T", "4.4", "Rule-based detectors versus Isolation Forest"))
    figure(ctx, "figs/rules_vs_ml.png", 5.2, "4.4", "Attacks found by the rule-based detectors and by the Isolation Forest")
    P_("The Isolation Forest found only the two scanning attacks, whose per-source statistics (many destination ports, many destination hosts) are far from normal. It missed DNS enumeration, FTP brute force and the rogue DHCP server, because these produce few packets that look statistically ordinary; it cannot see the ARP attack at all because ARP is not part of the feature set; and it missed the SYN flood because every spoofed source sends a single packet, so a per-source window carries no signal. It also raised false alarms on 2.5 % of normal windows. Protocol knowledge therefore matters more than generic statistics for this task, and the model is best used as a supplementary detector for unknown scanning behaviour.")
    heading(ctx, "4.7 Validation against the Objectives")
    table(ctx, ["Objective", "Evidence", "Status"], [
        ["O1 Reproducible lab and capture workflow", "Lab topology (Fig. 2.2); datasets regenerated identically by a seeded script; Appendix C gives live-capture commands", "Met"],
        ["O2 Baseline of normal traffic", "Table 4.1, Table 2.6 (normal maxima versus thresholds)", "Met"],
        ["O3 At least five attack classes", "Seven classes implemented (eight alert types)", "Met"],
        ["O4 Detection rate and latency", "7 of 7 detected; latency 0.00–1.45 s (Table 4.2)", "Met"],
        ["O5 Few false alarms, justified thresholds", "0 false alerts in normal, held-out and background traffic; sensitivity study (Table 4.3)", "Met on the test data"],
        ["O6 CLI and dashboard", "Fig. 3.1–3.3; dashboard runs on all sample captures", "Met"],
        ["O7 Machine-learning comparison", "Isolation Forest implemented and evaluated (Table 4.4); finds 2 of 7 attacks", "Partly met: implemented, limited accuracy"]],
        [2400, 4600, 1300], 10, ["l", "l", "c"], caption=("T", "4.5", "Validation of results against the project objectives"))
    P_("The only failed test case is TC-11, the slow scan, which is reported openly as a limitation rather than hidden by lowering the threshold.")
    heading(ctx, "4.8 Comparison with Established Tools")
    table(ctx, ["Aspect", "Snort / Zeek", "This system"], [
        ["Detection approach", "Signature rules and protocol analysers {r:roesch}{r:paxson}", "Protocol-aware thresholds plus optional Isolation Forest"],
        ["Port-scan logic", "Threshold or sequential-hypothesis testing, as in the Bro implementation {r:jung}", "Fixed 5-second distinct-port window"],
        ["Coverage", "Thousands of rules and many protocol analysers", "Seven attack classes"],
        ["Transparency for learning", "Large code base and rule language", "All detectors in 89 lines of Python"],
        ["Throughput", "Designed for live, high-rate links", f"Offline; 16,986 packets parsed in about 3.8 s on the development machine, rules evaluated in 0.03 s"],
        ["Evasion resistance", "Better, with multiple time scales and state", "Weak against slow scans (TC-11)"]],
        [1800, 3250, 3250], 10, caption=("T", "4.6", "Qualitative comparison with established intrusion-detection tools"))
    P_("The comparison is qualitative: no production tool was run on the project datasets. The prototype does not compete with these systems in coverage or speed; its value is that every decision can be traced to a single protocol fact.")

    # ---------- Chapter 5 ----------
    chapter(ctx, 5, "CONCLUSIONS")
    heading(ctx, "5.1 Key Conclusions")
    for t in ["Protocol-level rules derived from the course material (TCP handshake, ARP binding, DHCP offer, DNS query types, FTP reply codes) detect all seven attack classes studied, with latency between 0.00 and 1.45 s.",
              "Thresholds taken from a measured baseline gave no false alerts on 16,986 packets of held-out normal traffic and none in the background traffic of the attack captures; the margins are 6× to 12× for the volume-based rules.",
              "Several attacks (ARP spoofing, rogue DHCP, DNS zone transfer) change traffic volume very little, so they are visible only to detectors that understand the protocol and not to volume monitoring.",
              "An unsupervised Isolation Forest detected only the scanning attacks and produced false alarms on 2.5 % of normal windows; it complements but cannot replace protocol rules."]:
        B_(t)
    heading(ctx, "5.2 Extent to Which the Objectives Were Achieved")
    P_("Objectives O1 to O6 were achieved and O7 was achieved in part: the model was built and evaluated, but its detection ability is limited (Table 4.5).")
    heading(ctx, "5.3 Limitations")
    for t in ["The datasets were synthesised with Scapy to reproduce tool signatures. Real captures contain retransmissions, background noise and tool-specific details that this data does not model, so detection rates and the absence of false alerts are demonstrated on controlled data only.",
              "Thresholds are fixed. A slow scan (one probe every 1.5 s) evades the 5-second window rule, and lowering the threshold enough to catch it causes false alarms.",
              "The ARP detector must observe the legitimate binding before the attack, and the FTP detector relies on plain-text FTP; encrypted protocols (FTPS, DNS over HTTPS) are not inspected.",
              "Analysis is offline and in memory, only IPv4 is handled, and there is no automatic response."]:
        B_(t)
    heading(ctx, "5.4 Future Scope")
    for t in ["Capture traffic from the VirtualBox lab with the real tools (Appendix C) and from public datasets such as CICIDS2017 {r:sharafaldin}, and repeat the evaluation.",
              "Replace fixed windows with several time scales or sequential hypothesis testing {r:jung} to catch slow scans.",
              "Add IPv6 (neighbour-discovery spoofing), flow-based features and richer machine-learning models; run the detectors on a live capture stream.",
              "Connect alerts to automatic responses such as firewall rules, and to a central log or SIEM."]:
        B_(t)

    # ---------- References ----------
    P_("REFERENCES", "c", 12, bold=True, pbb=True, after=10)
    refp = []
    for i, k in enumerate(ctx.cites):
        q = ctx.doc.add_paragraph(); fmt(q, "j", 0, 6, 1.5)
        q.paragraph_format.left_indent = Pt(28); q.paragraph_format.first_line_indent = Pt(-28)
        add_runs(ctx, q, f"[{i+1}]\t" + REFS[k], 12)
        q.paragraph_format.tab_stops.add_tab_stop(Pt(28))
    assert not ctx.cites or len(set(ctx.cites)) == len(ctx.cites)

    # ---------- Appendices ----------
    p = P_("APPENDICES", "c", 12, bold=True, pbb=True, after=8)
    P_("This section contains supporting material that is not part of the main chapters: the contribution of the team members (Appendix A), the source code of the parser and detectors (Appendix B) and the commands for repeating the experiments on a live VirtualBox lab (Appendix C).")
    heading(ctx, "APPENDIX A: CONTRIBUTION OF TEAM MEMBERS", 2)
    table(ctx, ["Name", "Registration No.", "Major tasks / components", "Stages"], [
        ["[[NAME OF THE STUDENT 1]]", "[[24071A6701]]", "Lab design, dataset generator, packet parser, port-scan and ping-sweep detectors", "Planning, development, testing"],
        ["[[NAME OF THE STUDENT 2]]", "[[24071A6702]]", "ARP-spoofing, rogue-DHCP and DNS detectors, Isolation Forest, threshold study", "Development, evaluation"],
        ["[[NAME OF THE STUDENT 3]]", "[[24071A6703]]", "FTP and SYN-flood detectors, CLI, Streamlit dashboard, report", "Development, documentation"]],
        [1900, 1500, 3300, 1600], 10, caption=("T", "A.1", "Contribution of team members"))
    P_("The contribution of each member is clearly defined and mutually agreed upon by all members of the team.")
    heading(ctx, "APPENDIX B: SOURCE CODE", 2)
    P_("Project folder structure:")
    code(ctx, ["cbp_network_anomaly/", "  src/        gen_pcaps.py  parser.py  detectors.py  ml_iforest.py", "              run_detector.py  evaluate.py  make_figs.py", "  dashboard/  app.py", "  pcaps/      normal*.pcap, <attack>.pcap, combined.pcap, slowscan.pcap", "  results/    results.json, alerts_*.csv"], 8.5)
    P_("**B.1 parser.py** – converts a capture into a table of packets.", keep_next=True)
    code(ctx, open("src/parser.py").read().rstrip().split("\n"), 7)
    P_("**B.2 detectors.py** – the seven rule-based detectors.", keep_next=True)
    code(ctx, open("src/detectors.py").read().rstrip().split("\n"), 7)
    heading(ctx, "APPENDIX C: COMMANDS FOR THE LIVE LAB", 2)
    P_("To repeat the experiments with real traffic, start the capture on the monitor machine, run one attack from Kali against the victim, stop the capture and analyse the file. Run attacks only inside the isolated host-only network.")
    table(ctx, ["Attack", "Command on Kali (attacker)", "Detector expected to fire"], [
        ["Port scan", "nmap -sS -p 1-10000 192.168.56.10", "Port Scan"],
        ["Ping sweep", "nmap -sn 192.168.56.0/24", "Ping Sweep"],
        ["ARP spoofing", "arpspoof -i eth0 -t 192.168.56.20 192.168.56.1", "ARP Spoofing"],
        ["DNS enumeration, zone transfer", "dnsrecon -d lab.local -t brt\ndig axfr lab.local @192.168.56.10", "DNS Enumeration, DNS Zone Transfer"],
        ["FTP brute force", "hydra -l admin -P words.txt ftp://192.168.56.10", "FTP Brute Force"],
        ["SYN flood", "hping3 -S --flood --rand-source -p 80 192.168.56.10", "SYN Flood"],
        ["Rogue DHCP", "dnsmasq --interface=eth0 --dhcp-range=192.168.56.100,192.168.56.150 --no-daemon", "Rogue DHCP"]],
        [1900, 4300, 2100], 9.5, caption=("T", "C.1", "Attack commands for the live lab"))
    code(ctx, ["sudo tcpdump -i eth1 -w pcaps/live_portscan.pcap      # on the monitor machine", "python src/run_detector.py pcaps/live_portscan.pcap --csv results/live_alerts.csv"], 8.5)

    # ============ FRONT LISTS (moved before the section break) ============
    def to_front(fn):
        n = len(list(body.iterchildren())); fn(); new = list(body.iterchildren())[n - 1:]
        # the last element is the body sectPr; keep it in place
        for e in new[:-1] if new and new[-1].tag == qn("w:sectPr") else new: anchor.addprevious(e)
    def toc():
        rows = [["", "ABSTRACT", pg("abstract")], ["", "LIST OF TABLES", pg("lot")], ["", "LIST OF FIGURES", pg("lof")]]
        for n, t in [(1, "INTRODUCTION"), (2, "APPROACH / METHODOLOGY"), (3, "TEST CASES / OUTPUT"), (4, "RESULTS, ANALYSIS AND VALIDATION"), (5, "CONCLUSIONS")]:
            rows.append([f"CHAPTER {n}", t, pg(f"ch{n}")])
        rows += [["", "REFERENCES", pg("refs")], ["", "APPENDICES (A: Contribution of team members, B: Source code, C: Live-lab commands)", pg("apps")]]
        table(ctx, ["Chapter", "Title", "Page No."], rows, [1500, 5600, 1200], 12, ["l", "l", "r"])
    def abstract():
        front_heading(ctx, "ABSTRACT")
        P_("Modern networks carry many protocols, and attacks such as port scans, ARP spoofing and SYN floods leave recognisable but easily overlooked signatures in packet headers and protocol exchanges. This project presents a Network Traffic Analysis and Protocol-Level Anomaly Detection System that reads packet captures, reconstructs per-protocol behaviour and raises alerts when traffic deviates from protocol norms. The objectives were to design a reproducible lab, capture normal and attack traffic, implement detectors for the reconnaissance and attack techniques studied in the course, and evaluate them for detection rate, latency and false alarms. The system follows a capture–parse–detect–alert pipeline in Python: Scapy converts .pcap files into a per-packet table, seven rule-based detectors with sliding-window thresholds cover port scans, ping sweeps, ARP spoofing, DNS enumeration and zone transfer, FTP brute force, SYN flooding and rogue DHCP servers, and an Isolation Forest provides an exploratory unsupervised layer; results are shown through a command-line tool and a Streamlit dashboard. During the semester we built a labelled dataset of one baseline capture, a held-out normal capture, seven single-attack captures and one combined capture, generated with Scapy to reproduce the on-the-wire signatures of Nmap, hping3, Hydra, arpspoof, dig and a rogue DHCP server. All seven attack scenarios were detected, with detection latency between 0.00 and 1.45 seconds and no false alerts on 16,986 packets (ten minutes) of held-out normal traffic. The Isolation Forest found only the two scanning attacks, which shows that protocol-aware rules are more reliable than generic statistical features for this task, and a deliberately slow port scan evaded the five-second window rule, which exposes the main limitation of fixed thresholds. The work shows that the protocol knowledge from the Computer Networks and Ethical Hacking course can be turned into a working and explainable intrusion-detection prototype.", "j", 12, before=0, after=12)
    def lists():
        front_heading(ctx, "LIST OF TABLES")
        table(ctx, ["Table", "Title", "Page No."], [[f"Table {n}", t, pg(f"T:{n}")] for n, t in ctx.tables], [1200, 6000, 1100], 11, ["l", "l", "r"])
        front_heading(ctx, "LIST OF FIGURES")
        table(ctx, ["Figure", "Title", "Page No."], [[f"Fig. {n}", t, pg(f"F:{n}")] for n, t in ctx.figs], [1200, 6000, 1100], 11, ["l", "l", "r"])
    # figures and tables registries were filled during body build; split them
    tabs = [x for x in ctx.tables]; figs_ = [x for x in ctx.figs]
    # table() registers both under ctx.tables/ctx.figs depending on kind; figure() uses ctx.figs.
    to_front(toc); to_front(abstract)
    ctx.tables = [x for x in tabs]; ctx.figs = [x for x in figs_]
    to_front(lists)
    doc.core_properties.title = TITLE.title(); doc.core_properties.subject = SUBJECT
    doc.save(out); return ctx

def pdf_pages(docx):
    subprocess.run(["python", "/mnt/skills/public/docx/scripts/office/soffice.py", "--headless", "--convert-to", "pdf", "--outdir", os.path.dirname(docx) or ".", docx], capture_output=True)
    pdf = docx.replace(".docx", ".pdf"); txt = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True).stdout
    return [pg.split("\n") for pg in txt.split("\f")][:-1]

def page_map(pages, ctx):
    first = lambda i: next((l.strip() for l in pages[i] if l.strip()), "")
    ch1 = next(i for i in range(2, len(pages)) if first(i) == "CHAPTER 1")
    pr = lambda i: i if i < ch1 else i - ch1 + 1   # printed number (cover = pdf page 0; certificate = printed 1)
    pm = {}
    for i in range(1, len(pages)):
        f = first(i); pn = pr(i)
        if f == "ABSTRACT": pm["abstract"] = pn
        elif f == "LIST OF TABLES": pm["lot"] = pn
        elif f == "LIST OF FIGURES": pm["lof"] = pn
        elif re.fullmatch(r"CHAPTER [1-5]", f): pm["ch" + f[-1]] = pn
        elif f == "REFERENCES": pm["refs"] = pn
        elif f == "APPENDICES": pm["apps"] = pn
    for n, t in ctx.tables:
        for i in range(ch1, len(pages)):
            if any(l.strip().startswith(f"Table {n}:") for l in pages[i]): pm[f"T:{n}"] = pr(i); break
    for n, t in ctx.figs:
        for i in range(ch1, len(pages)):
            if any(l.strip().startswith(f"Fig. {n}:") for l in pages[i]): pm[f"F:{n}"] = pr(i); break
    return pm

if __name__ == "__main__":
    os.makedirs("report", exist_ok=True); out = "report/draft.docx"
    ctx = build({}, out); pages = pdf_pages(out); pm = page_map(pages, ctx)
    ctx = build(pm, out); pages2 = pdf_pages(out); pm2 = page_map(pages2, ctx)
    print("stable page map:", pm == pm2, "| pdf pages:", len(pages2)); json.dump(pm2, open("report/pagemap.json", "w"), indent=1)
    missing = [k for k in pm2 if False]
    if pm != pm2: ctx = build(pm2, out); print("rebuilt with second map")
