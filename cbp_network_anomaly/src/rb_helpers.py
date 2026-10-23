"""Helpers used by build_report.py (fills the VNRVJIET CBP template)."""
import re, copy
from docx.shared import Pt, Inches, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml
from docx.text.paragraph import Paragraph

TEXTW = 8306  # A4 with 1800-twip side margins
ALIGN = {"l": AL.LEFT, "c": AL.CENTER, "j": AL.JUSTIFY, "r": AL.RIGHT}

REFS = {
 "forouzan": "B. A. Forouzan, *Data Communications and Networking*, 4th ed. New York, NY, USA: McGraw-Hill, 2006.",
 "kurose": "J. F. Kurose and K. W. Ross, *Computer Networking: A Top-Down Approach*, 6th ed. Boston, MA, USA: Pearson, 2012.",
 "engebretson": "P. Engebretson, *The Basics of Hacking and Penetration Testing*, 2nd ed. Waltham, MA, USA: Syngress (Elsevier), 2013.",
 "lyon": "G. F. Lyon, *Nmap Network Scanning: The Official Nmap Project Guide to Network Discovery and Security Scanning*. Sunnyvale, CA, USA: Insecure.Com LLC, 2009.",
 "rfc793": "J. Postel, “Transmission Control Protocol,” IETF, RFC 793, Sep. 1981.",
 "rfc792": "J. Postel, “Internet Control Message Protocol,” IETF, RFC 792, Sep. 1981.",
 "rfc826": "D. C. Plummer, “An Ethernet Address Resolution Protocol,” IETF, RFC 826, Nov. 1982.",
 "rfc2131": "R. Droms, “Dynamic Host Configuration Protocol,” IETF, RFC 2131, Mar. 1997.",
 "rfc1035": "P. Mockapetris, “Domain names - implementation and specification,” IETF, RFC 1035, Nov. 1987.",
 "rfc5936": "E. Lewis and A. Hoenes, “DNS Zone Transfer Protocol (AXFR),” IETF, RFC 5936, Jun. 2010.",
 "rfc959": "J. Postel and J. Reynolds, “File Transfer Protocol (FTP),” IETF, RFC 959, Oct. 1985.",
 "rfc4987": "W. Eddy, “TCP SYN flooding attacks and common mitigations,” IETF, RFC 4987, Aug. 2007.",
 "scapy": "P. Biondi and the Scapy community, “Scapy: Packet crafting for Python.” [Online]. Available: https://scapy.net",
 "wireshark": "“Wireshark: Network protocol analyzer,” Wireshark Foundation. [Online]. Available: https://www.wireshark.org",
 "pandas": "W. McKinney, “Data structures for statistical computing in Python,” in *Proc. 9th Python in Science Conf. (SciPy)*, Austin, TX, USA, Jun. 2010, pp. 56–61.",
 "sklearn": "F. Pedregosa *et al.*, “Scikit-learn: Machine learning in Python,” *J. Mach. Learn. Res.*, vol. 12, pp. 2825–2830, Oct. 2011.",
 "liu": "F. T. Liu, K. M. Ting, and Z.-H. Zhou, “Isolation forest,” in *Proc. 8th IEEE Int. Conf. Data Mining (ICDM)*, Pisa, Italy, Dec. 2008, pp. 413–422.",
 "streamlit": "“Streamlit: A faster way to build and share data apps.” [Online]. Available: https://streamlit.io",
 "roesch": "M. Roesch, “Snort: Lightweight intrusion detection for networks,” in *Proc. 13th USENIX Large Installation Syst. Admin. Conf. (LISA)*, Seattle, WA, USA, Nov. 1999, pp. 229–238.",
 "paxson": "V. Paxson, “Bro: A system for detecting network intruders in real-time,” *Comput. Netw.*, vol. 31, no. 23–24, pp. 2435–2463, Dec. 1999.",
 "jung": "S. Jung, V. Paxson, A. W. Berger, and H. Balakrishnan, “Fast portscan detection using sequential hypothesis testing,” in *Proc. IEEE Symp. Security and Privacy*, Oakland, CA, USA, May 2004, pp. 211–225.",
 "sharafaldin": "I. Sharafaldin, A. H. Lashkari, and A. A. Ghorbani, “Toward generating a new intrusion detection dataset and intrusion traffic characterization,” in *Proc. 4th Int. Conf. Inf. Syst. Security Privacy (ICISSP)*, Funchal, Portugal, Jan. 2018, pp. 108–116.",
}

class Ctx:
    """Holds document, citation numbering and caption registries."""
    def __init__(self, doc):
        self.doc = doc; self.cites = []; self.tables = []; self.figs = []
    def cite(self, text):
        def sub(m):
            k = m.group(1)
            if k not in self.cites: self.cites.append(k)
            return f"[{self.cites.index(k)+1}]"
        return re.sub(r"\{r:(\w+)\}", sub, text)

TOK = re.compile(r"(\*\*.+?\*\*|\[\[.+?\]\]|`.+?`|\*[^*\s][^*]*?\*)")

def setfont(run, name):
    rPr = run._r.get_or_add_rPr(); rf = rPr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rPr.insert(0, rf)
    for a in ("ascii", "hAnsi", "eastAsia", "cs"): rf.set(qn("w:" + a), name)

def add_runs(ctx, p, text, size=12, bold=False, italic=False, color=None):
    text = ctx.cite(text)
    for part in TOK.split(text):
        if not part: continue
        b, i, c, f, t = bold, italic, color, None, part
        if part.startswith("**"): t, b = part[2:-2], True
        elif part.startswith("[["): t, b, c = part[2:-2], True, "FF0000"
        elif part.startswith("`"): t, f = part[1:-1], "Courier New"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2: t, i = part[1:-1], True
        r = p.add_run(t); r.font.size = Pt(size if not f else size - 1)
        if b: r.bold = True
        if i: r.italic = True
        if c: r.font.color.rgb = RGBColor.from_string(c)
        if f: setfont(r, f)

def fmt(p, align="j", before=0, after=6, line=1.5, keep_next=False, pbb=False, keep_lines=False):
    pf = p.paragraph_format; pf.alignment = ALIGN[align]; pf.space_before = Pt(before); pf.space_after = Pt(after)
    pf.line_spacing = line; pf.keep_with_next = keep_next; pf.page_break_before = pbb; pf.keep_together = keep_lines

def P(ctx, text, align="j", size=12, bold=False, italic=False, before=0, after=6, line=1.5, keep_next=False, pbb=False, color=None):
    p = ctx.doc.add_paragraph(); fmt(p, align, before, after, line, keep_next, pbb)
    add_runs(ctx, p, text, size, bold, italic, color); return p

def spacer(ctx, pts=4):
    p = ctx.doc.add_paragraph(); fmt(p, "l", 0, 0, 1.0); r = p.add_run(""); r.font.size = Pt(pts)
    p.paragraph_format.line_spacing = Pt(pts + 2); return p

def bullet(ctx, text, level=0, size=12, after=3):
    p = P(ctx, text, "j", size, after=after)
    numPr = p._p.get_or_add_pPr().get_or_add_numPr(); numPr.get_or_add_ilvl().val = level; numPr.get_or_add_numId().val = 90
    return p

def heading(ctx, text, level=2):
    p = ctx.doc.add_paragraph(style=f"Heading {level}")
    fmt(p, "l", 10 if level == 2 else 6, 4, 1.5, keep_next=True)
    r = p.add_run(text); setfont(r, "Times New Roman"); r.font.size = Pt(12); r.font.color.rgb = RGBColor(0, 0, 0); r.bold = True
    if level == 3: r.italic = True
    return p

def chapter(ctx, num, title):
    for k, txt in enumerate([f"CHAPTER {num}", title]):
        p = ctx.doc.add_paragraph(style="Heading 1"); fmt(p, "c", 0, 0 if k == 0 else 14, 1.5, keep_next=True, pbb=(k == 0))
        r = p.add_run(txt); setfont(r, "Times New Roman"); r.font.size = Pt(12); r.font.color.rgb = RGBColor(0, 0, 0)

def front_heading(ctx, text, pbb=True):
    p = ctx.doc.add_paragraph(); fmt(p, "c", 0, 12, 1.5, keep_next=True, pbb=pbb)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = RGBColor(0xC0, 0, 0); return p

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def table(ctx, headers, rows, widths, size=10, aligns=None, caption=None, bold_first_col=False):
    assert sum(widths) <= TEXTW + 5, sum(widths)
    if caption:
        kind, num, title = caption
        (ctx.tables if kind == "T" else ctx.figs).append((num, title))
        P(ctx, f"Table {num}: {title}", "c", 11, bold=True, before=6, after=4, line=1.0, keep_next=True)
    ncol = len(headers); t = ctx.doc.add_table(rows=1 + len(rows), cols=ncol); t.style = ctx.doc.styles["Table Grid"]
    t.autofit = False; tblPr = t._tbl.tblPr
    tw = tblPr.find(qn("w:tblW"))
    if tw is None: tw = OxmlElement("w:tblW"); tblPr.append(tw)
    tw.set(qn("w:w"), str(sum(widths))); tw.set(qn("w:type"), "dxa")
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths): gc.set(qn("w:w"), str(w))
    aligns = aligns or ["l"] * ncol
    for ri, row in enumerate([headers] + rows):
        tr = t.rows[ri]; trPr = tr._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        if ri == 0: th = OxmlElement("w:tblHeader"); trPr.append(th)
        for ci, val in enumerate(row):
            cell = tr.cells[ci]; cell.width = Twips(widths[ci])
            if ri == 0: shade(cell, "D9E2F3")
            lines = str(val).split("\n")
            for li, ln in enumerate(lines):
                p = cell.paragraphs[0] if li == 0 else cell.add_paragraph()
                fmt(p, "c" if ri == 0 else aligns[ci], 2, 2, 1.0)
                add_runs(ctx, p, ln, size, bold=(ri == 0 or (bold_first_col and ci == 0)))
    if len(rows) <= 12:
        for tr in t.rows[:-1]:
            for c in tr.cells:
                for pp in c.paragraphs: pp.paragraph_format.keep_with_next = True
    spacer(ctx, 6); return t

def figure(ctx, path, width, num, title):
    ctx.figs.append((num, title)); ctx.doc.add_picture(path, width=Inches(width))
    p = ctx.doc.paragraphs[-1]; fmt(p, "c", 6, 2, 1.0, keep_next=True)
    P(ctx, f"Fig. {num}: {title}", "c", 11, bold=True, before=2, after=10, line=1.0)

def code(ctx, lines, size=8):
    import textwrap
    wrapped = []
    cap = int(88 * 8 / size) if size < 8 else 88
    for ln in lines:
        if len(ln) <= cap: wrapped.append(ln); continue
        ind = len(ln) - len(ln.lstrip())
        wrapped += textwrap.wrap(ln, cap, subsequent_indent=" " * (ind + 4), break_long_words=True, break_on_hyphens=False) or [""]
    lines = wrapped
    for i, ln in enumerate(lines):
        p = ctx.doc.add_paragraph(); fmt(p, "l", 0, 0, 1.0)
        r = p.add_run(ln if ln else " "); r.font.size = Pt(size); setfont(r, "Courier New")
        pPr = p._p.get_or_add_pPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F2F2F2")
        succ = [qn("w:" + n) for n in ("tabs","suppressAutoHyphens","kinsoku","wordWrap","overflowPunct","topLinePunct","autoSpaceDE","autoSpaceDN","bidi","adjustRightInd","snapToGrid","spacing","ind","contextualSpacing","mirrorIndents","suppressOverlap","jc","textDirection","textAlignment","textboxTightWrap","outlineLvl","divId","cnfStyle","rPr","sectPr","pPrChange")]
        nxt = next((c for c in pPr if c.tag in succ), None)
        (nxt.addprevious(shd) if nxt is not None else pPr.append(shd))
    spacer(ctx, 6)

def replace_in_para(p_el, old, new, color="000000"):
    ts = p_el.findall(".//" + qn("w:t")); full = "".join(t.text or "" for t in ts)
    i = full.find(old); assert i >= 0, (old, full[:80]); j = i + len(old); pos = 0; first = True
    for t in ts:
        s = t.text or ""; a, b = pos, pos + len(s); pos = b
        if b <= i or a >= j: continue
        pre = s[:i - a] if a < i else ""; suf = s[j - a:] if b > j else ""
        if first:
            t.text = pre + new + suf; first = False
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            rPr = t.getparent().find(qn("w:rPr"))
            if rPr is not None:
                c = rPr.find(qn("w:color"))
                if c is not None: c.set(qn("w:val"), color)
        else:
            t.text = suf
